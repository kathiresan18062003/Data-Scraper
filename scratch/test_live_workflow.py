import os
import sys
import openpyxl
from pathlib import Path

# Add project root to path
sys.path.insert(0, str(Path(__file__).parent.parent))

from fastapi.testclient import TestClient
from backend.app.main import app
from backend.app.services import excel_service

client = TestClient(app)

def test_live_workflow():
    print("--- 1. Creating Representative Excel Workbook ---")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalog"
    ws.append(["S.No", "MPN", "MPN1", "Manufacturer", "Product Name", "UPC", "Source URL"])
    ws.append(["1", "140-13303", "140-13303-R", "Wilwood", "Wilwood Front Brake Caliper Kit", "889545028441", "https://wilwood.com/calipers/140-13303-R"])
    ws.append(["2", "20140", "20140", "Moroso", "Moroso Street/Strip Oil Pan", "084663201403", "https://moroso.com/oil-pan-20140"])
    
    test_file = Path("data/uploads/live_test_catalog.xlsx")
    test_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(test_file)
    print(f"Saved test workbook to: {test_file}")

    print("\n--- 2. Uploading Workbook via API ---")
    with open(test_file, "rb") as f:
        res = client.post("/api/jobs/upload", files={"file": ("live_test_catalog.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert res.status_code == 200, f"Upload failed: {res.text}"
    upload_data = res.json()
    file_token = upload_data["file_token"]
    sheet = upload_data["sheets"][0]
    print(f"Upload successful. File Token: {file_token}, Sheet: {sheet['name']}, Detected Columns: {sheet['detected_columns']}")

    print("\n--- 3. Starting Job with MPN1 Primary Mapping ---")
    col_mapping = {
        "sno": "S.No",
        "mpn": "MPN",
        "mpn1": "MPN1",
        "mfr": "Manufacturer",
        "title": "Product Name",
        "upc": "UPC",
        "url": "Source URL"
    }
    start_payload = {
        "file_token": file_token,
        "sheet_name": sheet["name"],
        "column_mapping": col_mapping,
        "limit": 2,
        "timeout": 15,
        "retries": 1,
        "use_gemini": False,
        "save_every": 1
    }
    start_res = client.post("/api/jobs", json=start_payload)
    assert start_res.status_code == 200, f"Job start failed: {start_res.text}"
    job = start_res.json()
    job_id = job["id"]
    print(f"Job created successfully. ID: {job_id}, Status: {job['status']}, Total Products: {job['total_products']}")

    print("\n--- 4. Executing Step 1 (Product 1) ---")
    from backend.app.services.job_runner import process_job_step
    more1 = process_job_step(job_id)
    print(f"Step 1 executed. More rows remaining: {more1}")

    # Check live checkpoint
    cp = excel_service.load_checkpoint(job_id)
    assert cp is not None, "Checkpoint file was not created!"
    print(f"Checkpoint verified: {len(cp.get('saved_rows', {}))} product(s) saved in checkpoint.")
    
    # Check live Excel organized output
    org_path = Path("output") / f"{job_id}_Final_Organized_Output.xlsx"
    assert org_path.exists(), f"Organized output file not found: {org_path}"
    wb_org = openpyxl.load_workbook(org_path)
    ws_org = wb_org["Organized Output"]
    row2_vals = [cell.value for cell in ws_org[2]]
    print(f"Row 2 Organized Output: MPN={row2_vals[1]}, MFR={row2_vals[2]}, Status={row2_vals[4]}, Score={row2_vals[5]}")
    wb_org.close()

    print("\n--- 5. Testing Interruption / Resumption Safety ---")
    # Pausing job
    pause_res = client.post(f"/api/jobs/{job_id}/pause")
    assert pause_res.status_code == 200
    print("Job paused successfully.")

    # Resuming job
    resume_res = client.post(f"/api/jobs/{job_id}/resume")
    assert resume_res.status_code == 200
    print("Job resumed successfully.")

    print("\n--- 6. Executing Step 2 (Product 2) ---")
    more2 = process_job_step(job_id)
    print(f"Step 2 executed. More rows remaining: {more2}")

    # Final step completion
    more3 = process_job_step(job_id)
    print(f"Step 3 executed (completion check). More rows remaining: {more3}")

    job_final_res = client.get(f"/api/jobs/{job_id}")
    job_final = job_final_res.json()
    print(f"Final Job Status: {job_final['status']}, Processed: {job_final['processed_products']}/{job_final['total_products']}")
    print(f"Verified OK: {job_final['successful_products']}, Review: {job_final['review_products']}, Failed: {job_final['failed_products']}")

    # Inspect generated Scraped & Review workbooks
    scraped_path = Path("output") / f"{job_id}_Data_collection_scraped.xlsx"
    assert scraped_path.exists(), "Scraped workbook not created!"
    wb_scraped = openpyxl.load_workbook(scraped_path)
    assert "Catalog" in wb_scraped.sheetnames, "Original sheet structure not preserved!"
    wb_scraped.close()

    review_path = Path("output") / f"{job_id}_Scrape_Review_Report.xlsx"
    assert review_path.exists(), "Review workbook not created!"
    wb_review = openpyxl.load_workbook(review_path)
    assert "Summary" in wb_review.sheetnames and "Row Audit" in wb_review.sheetnames, "Review tabs not created!"
    wb_review.close()

    print("\n=== LIVE WORKFLOW TEST PASSED 100% SUCCESSFULLY ===")

if __name__ == "__main__":
    test_live_workflow()

