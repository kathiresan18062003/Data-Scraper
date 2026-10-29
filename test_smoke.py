from backend.data_validator import validate, strip_line_code
from backend.scraper import parse
import openpyxl
from backend import excel_handler as xh, export_handler as eh

def test_line_code():
    assert strip_line_code("WLD140-13303-R") == "140-13303-R"

def test_validate_exact():
    p = {"mpn": "140-13303-R", "product_title": "Wilwood Caliper", "manufacturer": "Wilwood", "text": ""}
    assert validate("WLD140-13303-R", "Wilwood", "https://autoplicity.com/p/140-13303-r", p, {"autoplicity.com"})["is_valid"]

def test_reject_search_page():
    assert not validate("X1234", "A", "https://x.com/search?q=X1234", {"text": "X1234"})["is_valid"]

def test_parse_jsonld():
    h = '<script type="application/ld+json">{"@type":"Product","name":"N","sku":"AB123","brand":{"name":"B"},"image":"https://x.com/a.jpg","description":"D"}</script>'
    r = parse(h, "https://x.com/p")
    assert r["mpn"] == "AB123" and r["image_urls"] == ["https://x.com/a.jpg"]

def test_export_preserves_columns_and_order(tmp_path):
    p = tmp_path / "in.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["MPN", "Manufacturer"]); ws.append(["A1234", "Acme"]); ws.append(["B5678", "Acme"])
    wb.save(p)
    ws = openpyxl.load_workbook(p)["Data"]
    rows = xh.read_rows(ws, xh.detect_columns(ws))
    res = {2: {"status": "SUCCESS", "sources": ["https://x.com/a1234"], "overview": "Ov", "features": ["f1"],
               "specifications": [{"name": "K", "value": "V"}], "image": "https://x.com/i.jpg", "google": "g"},
           3: {"status": "FAILED", "sources": [], "reason": "none", "google": "g"}}
    out, _ = eh.export(p, "Data", rows, res)
    o = openpyxl.load_workbook(out)["Data"]
    h = [c.value for c in o[1]]
    assert h[:2] == ["MPN", "Manufacturer"] and o["A2"].value == "A1234" and o["A3"].value == "B5678"
    assert o.cell(2, h.index("Source 1") + 1).value == "https://x.com/a1234"
    assert o.cell(3, h.index("DataNotAvailable") + 1).value == "Y"

    org = openpyxl.load_workbook(eh.settings.OUTPUT_DIR / eh.ORGANIZED)["Organized Output"]
    oh = [c.value for c in org[1]]
    assert oh[:3] == ["Row", "MPN", "Manufacturer"] and org.max_row == 3
    assert org.cell(2, oh.index("Status") + 1).value == "SUCCESS"
    assert "K: V" in org.cell(2, oh.index("Specifications") + 1).value


def test_export_called_every_n_rows_plus_at_end(monkeypatch, tmp_path):
    from backend import pipeline
    rows = [
        {"row": 2, "mpn": "P1", "mfr": "M1"},
        {"row": 3, "mpn": "P2", "mfr": "M1"},
        {"row": 4, "mpn": "P3", "mfr": "M1"},
        {"row": 5, "mpn": "P4", "mfr": "M1"},
        {"row": 6, "mpn": "P5", "mfr": "M1"},
    ]
    results = {}
    ckpt = tmp_path / "ckpt.json"

    monkeypatch.setattr(pipeline, "process_row", lambda r, *a, **k: {
        "status": "SUCCESS", "tier": "1", "sources": ["https://x.com"], "score": 90,
        "confidence": "HIGH", "product_title": f"Prod {r['mpn']}"
    })

    export_calls = []
    def mock_on_export(r, res, is_final=False):
        export_calls.append((len(res), is_final))

    cfg = {"save_every": 2, "on_export": mock_on_export}
    pipeline.run_batch(rows, results, ckpt, cfg)

    # 5 rows processed, save_every = 2
    # Triggers at row 2, row 4, and once at the end
    assert len(export_calls) == 3
    assert export_calls[0] == (2, False)
    assert export_calls[1] == (4, False)
    assert export_calls[2] == (5, True)


def test_partial_export_contains_processed_and_pending_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(eh.settings, "OUTPUT_DIR", tmp_path)
    p = tmp_path / "input.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Products"
    ws.append(["MPN", "Manufacturer"])
    ws.append(["PART1", "BrandA"])
    ws.append(["PART2", "BrandA"])
    ws.append(["PART3", "BrandA"])
    wb.save(p)

    rows = [
        {"row": 2, "mpn": "PART1", "mfr": "BrandA"},
        {"row": 3, "mpn": "PART2", "mfr": "BrandA"},
        {"row": 4, "mpn": "PART3", "mfr": "BrandA"},
    ]
    results = {
        2: {"status": "SUCCESS", "sources": ["https://brand.com/part1"], "overview": "Overview 1",
            "features": ["F1"], "specifications": [{"name": "S1", "value": "V1"}],
            "image": "https://brand.com/img.jpg", "google": "https://google.com/q1", "notes": []}
    }

    scraped_path, review_path = eh.export(p, "Products", rows, results, is_final=False)

    # Re-open and verify organized output
    org_wb = openpyxl.load_workbook(tmp_path / eh.ORGANIZED)
    org_ws = org_wb["Organized Output"]
    assert org_ws.max_row == 4
    headers = [c.value for c in org_ws[1]]
    status_col = headers.index("Status") + 1

    assert org_ws.cell(2, status_col).value == "SUCCESS"
    assert org_ws.cell(3, status_col).value == "PENDING"
    assert org_ws.cell(4, status_col).value == "PENDING"

    # Re-open and verify review output
    rev_wb = openpyxl.load_workbook(review_path)
    rev_audit = rev_wb["Row Audit"]
    audit_headers = [c.value for c in rev_audit[1]]
    audit_status_col = audit_headers.index("Status") + 1

    assert rev_audit.cell(2, audit_status_col).value == "SUCCESS"
    assert rev_audit.cell(3, audit_status_col).value == "PENDING"
    assert rev_audit.cell(4, audit_status_col).value == "PENDING"


def test_scraped_preserves_original_columns_and_order(tmp_path, monkeypatch):
    monkeypatch.setattr(eh.settings, "OUTPUT_DIR", tmp_path)
    p = tmp_path / "custom_input.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "CustomSheet"
    ws.append(["CustomID", "MPN", "SupplierCode", "Manufacturer", "InternalNotes"])
    ws.append(["ID-001", "MPN-101", "SUP-A", "Mfr1", "Keep note 1"])
    ws.append(["ID-002", "MPN-102", "SUP-B", "Mfr2", "Keep note 2"])
    ws.append(["ID-003", "MPN-103", "SUP-C", "Mfr3", "Keep note 3"])
    wb.save(p)

    rows = [
        {"row": 2, "mpn": "MPN-101", "mfr": "Mfr1"},
        {"row": 3, "mpn": "MPN-102", "mfr": "Mfr2"},
        {"row": 4, "mpn": "MPN-103", "mfr": "Mfr3"},
    ]
    results = {
        2: {"status": "SUCCESS", "sources": ["https://mfr1.com/p"], "overview": "Ov",
            "features": ["F"], "specifications": [], "image": "", "google": "g1"},
        3: {"status": "FAILED", "sources": [], "reason": "Not found", "google": "g2"}
    }

    scraped_path, _ = eh.export(p, "CustomSheet", rows, results, is_final=False)

    # Re-open and verify Data_collection_scraped.xlsx
    out_wb = openpyxl.load_workbook(scraped_path)
    out_ws = out_wb["CustomSheet"]

    out_headers = [c.value for c in out_ws[1]]
    assert out_headers[:5] == ["CustomID", "MPN", "SupplierCode", "Manufacturer", "InternalNotes"]

    assert out_ws.cell(2, 1).value == "ID-001"
    assert out_ws.cell(2, 2).value == "MPN-101"
    assert out_ws.cell(2, 5).value == "Keep note 1"

    assert out_ws.cell(3, 1).value == "ID-002"
    assert out_ws.cell(3, 2).value == "MPN-102"
    assert out_ws.cell(3, 5).value == "Keep note 2"

    assert out_ws.cell(4, 1).value == "ID-003"
    assert out_ws.cell(4, 2).value == "MPN-103"
    assert out_ws.cell(4, 5).value == "Keep note 3"

    comp_col = out_headers.index("Completed") + 1
    assert out_ws.cell(2, comp_col).value == "Y"
    assert out_ws.cell(3, comp_col).value == "N"
    assert out_ws.cell(4, comp_col).value is None


def test_failing_export_does_not_stop_batch(monkeypatch, tmp_path):
    from backend import pipeline
    rows = [
        {"row": 2, "mpn": "P1", "mfr": "M1"},
        {"row": 3, "mpn": "P2", "mfr": "M1"},
        {"row": 4, "mpn": "P3", "mfr": "M1"},
    ]
    results = {}
    ckpt = tmp_path / "ckpt.json"

    monkeypatch.setattr(pipeline, "process_row", lambda r, *a, **k: {
        "status": "SUCCESS", "tier": "1", "sources": [], "score": 80,
        "confidence": "HIGH", "product_title": f"Title {r['mpn']}"
    })

    def buggy_export(r, res, **kw):
        raise PermissionError("Simulated write lock error on export file")

    logs = []
    cfg = {"save_every": 1, "on_export": buggy_export}
    res = pipeline.run_batch(rows, results, ckpt, cfg, on_log=logs.append)

    assert len(res) == 3
    assert set(res.keys()) == {2, 3, 4}
    assert any("Export error" in m for m in logs)


def test_resume_after_interruption_no_duplicates(tmp_path, monkeypatch):
    monkeypatch.setattr(eh.settings, "OUTPUT_DIR", tmp_path)
    from backend import pipeline
    import pytest

    p = tmp_path / "in.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["MPN", "Manufacturer"])
    ws.append(["P100", "BrandX"])
    ws.append(["P200", "BrandX"])
    ws.append(["P300", "BrandX"])
    ws.append(["P400", "BrandX"])
    wb.save(p)

    rows = [
        {"row": 2, "mpn": "P100", "mfr": "BrandX"},
        {"row": 3, "mpn": "P200", "mfr": "BrandX"},
        {"row": 4, "mpn": "P300", "mfr": "BrandX"},
        {"row": 5, "mpn": "P400", "mfr": "BrandX"},
    ]
    ckpt = tmp_path / "ckpt.json"

    call_count = [0]
    def process_with_interruption(r, *a, **k):
        call_count[0] += 1
        if call_count[0] == 3:
            raise KeyboardInterrupt("Simulated interruption during batch")
        return {
            "status": "SUCCESS", "tier": "1", "sources": [f"https://x.com/{r['mpn']}"],
            "score": 95, "confidence": "HIGH", "product_title": f"Product {r['mpn']}",
            "overview": "Overview text", "features": ["Feature A"],
            "specifications": [{"name": "Spec", "value": "Val"}], "image": "", "notes": []
        }

    monkeypatch.setattr(pipeline, "process_row", process_with_interruption)

    results = {}
    cfg = {"save_every": 2, "on_export": lambda r, res, **kw: eh.export(p, "Data", r, res, **kw)}

    with pytest.raises(KeyboardInterrupt):
        pipeline.run_batch(rows, results, ckpt, cfg)

    assert len(results) == 2
    loaded_ckpt = xh.load_ckpt(ckpt)
    assert len(loaded_ckpt) == 2

    # Partial export was created
    org_wb = openpyxl.load_workbook(tmp_path / eh.ORGANIZED)
    assert org_wb["Organized Output"].max_row == 5
    status_vals_run1 = [org_wb["Organized Output"].cell(r, 5).value for r in range(2, 6)]
    assert status_vals_run1 == ["SUCCESS", "SUCCESS", "PENDING", "PENDING"]

    # Resume without interruption
    monkeypatch.setattr(pipeline, "process_row", lambda r, *a, **k: {
        "status": "SUCCESS", "tier": "1", "sources": [f"https://x.com/{r['mpn']}"],
        "score": 95, "confidence": "HIGH", "product_title": f"Product {r['mpn']}",
        "overview": "Overview text", "features": ["Feature A"],
        "specifications": [{"name": "Spec", "value": "Val"}], "image": "", "notes": []
    })

    results2 = xh.load_ckpt(ckpt)
    pipeline.run_batch(rows, results2, ckpt, cfg)

    assert len(results2) == 4
    final_org = openpyxl.load_workbook(tmp_path / eh.ORGANIZED)["Organized Output"]
    assert final_org.max_row == 5
    mpn_vals = [final_org.cell(r, 2).value for r in range(2, 6)]
    assert mpn_vals == ["P100", "P200", "P300", "P400"]
    status_vals_final = [final_org.cell(r, 5).value for r in range(2, 6)]
    assert status_vals_final == ["SUCCESS", "SUCCESS", "SUCCESS", "SUCCESS"]


def test_periodic_save_option_off(monkeypatch, tmp_path):
    from backend import pipeline
    rows = [
        {"row": 2, "mpn": "P1", "mfr": "M1"},
        {"row": 3, "mpn": "P2", "mfr": "M1"},
        {"row": 4, "mpn": "P3", "mfr": "M1"},
    ]
    results = {}
    ckpt = tmp_path / "ckpt.json"

    monkeypatch.setattr(pipeline, "process_row", lambda r, *a, **k: {
        "status": "SUCCESS", "tier": "1", "sources": [], "score": 85,
        "confidence": "HIGH", "product_title": f"P {r['mpn']}"
    })

    calls = []
    def record_export(r, res, is_final=False):
        calls.append((len(res), is_final))

    # Option to turn off: save_every=0
    cfg = {"save_every": 0, "on_export": record_export}
    pipeline.run_batch(rows, results, ckpt, cfg)

    # When turned off, periodic export is not called, called only once at the end
    assert len(calls) == 1
    assert calls[0] == (3, True)


def test_history_json_not_duplicated_and_marked_partial(tmp_path, monkeypatch):
    import json
    monkeypatch.setattr(eh.settings, "OUTPUT_DIR", tmp_path)
    p = tmp_path / "in.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["MPN", "Manufacturer"])
    ws.append(["A1", "Brand"])
    ws.append(["A2", "Brand"])
    ws.append(["A3", "Brand"])
    wb.save(p)

    rows = [
        {"row": 2, "mpn": "A1", "mfr": "Brand"},
        {"row": 3, "mpn": "A2", "mfr": "Brand"},
        {"row": 4, "mpn": "A3", "mfr": "Brand"},
    ]
    res1 = {2: {"status": "SUCCESS", "sources": []}}
    # Periodic partial export 1: is_final=False
    eh.export(p, "Data", rows, res1, is_final=False)

    hist_file = tmp_path / "history.json"
    # No history written for periodic partial export
    assert not hist_file.exists()

    res2 = {2: {"status": "SUCCESS", "sources": []}, 3: {"status": "FAILED", "sources": []}}
    # Periodic partial export 2: is_final=False
    eh.export(p, "Data", rows, res2, is_final=False)
    assert not hist_file.exists()

    # Interrupted end of run: is_final=True but only 2 of 3 rows completed
    eh.export(p, "Data", rows, res2, is_final=True)
    assert hist_file.exists()
    h = json.loads(hist_file.read_text())
    assert len(h) == 1
    assert h[0]["status"] == "PARTIAL"
    assert h[0]["total"] == 3
    assert h[0]["success"] == 1
    assert h[0]["failed"] == 1

    # Final completed run with all 3 rows
    res3 = {2: {"status": "SUCCESS", "sources": []}, 3: {"status": "FAILED", "sources": []}, 4: {"status": "SUCCESS", "sources": []}}
    eh.export(p, "Data", rows, res3, is_final=True)
    h = json.loads(hist_file.read_text())
    assert len(h) == 2
    assert h[1]["status"] == "COMPLETE"


def test_immediate_per_product_saving(monkeypatch, tmp_path):
    from backend import pipeline
    rows = [
        {"row": 2, "mpn": "IMM1", "mfr": "Brand"},
        {"row": 3, "mpn": "IMM2", "mfr": "Brand"},
        {"row": 4, "mpn": "IMM3", "mfr": "Brand"},
    ]
    results = {}
    ckpt = tmp_path / "ckpt.json"

    monkeypatch.setattr(pipeline, "process_row", lambda r, *a, **k: {
        "status": "SUCCESS", "tier": "1", "sources": [], "score": 90,
        "confidence": "HIGH", "product_title": f"P {r['mpn']}"
    })

    export_log = []
    def log_export(r, res, is_final=False):
        export_log.append((len(res), is_final))

    # Default cfg has no explicit save_every -> defaults to 1 (immediate per product)
    cfg = {"on_export": log_export}
    pipeline.run_batch(rows, results, ckpt, cfg)

    # For 3 rows: called after row 1 (len 1), after row 2 (len 2), after row 3 (len 3), plus once at the end (len 3)
    assert len(export_log) == 4
    assert export_log[0] == (1, False)
    assert export_log[1] == (2, False)
    assert export_log[2] == (3, False)
    assert export_log[3] == (3, True)


def test_concurrent_export_lock_safety(tmp_path, monkeypatch):
    import concurrent.futures
    monkeypatch.setattr(eh.settings, "OUTPUT_DIR", tmp_path)
    p = tmp_path / "in.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Data"
    ws.append(["MPN", "Manufacturer"])
    for i in range(1, 11):
        ws.append([f"MPN-{i}", "Brand"])
    wb.save(p)

    rows = [{"row": i + 1, "mpn": f"MPN-{i}", "mfr": "Brand"} for i in range(1, 11)]

    # Multiple worker threads calling eh.export concurrently
    def worker(worker_id):
        res = {r["row"]: {"status": "SUCCESS", "sources": [], "product_title": f"Title {worker_id}"}
               for r in rows[:worker_id + 1]}
        eh.export(p, "Data", rows, res, is_final=False)

    with concurrent.futures.ThreadPoolExecutor(max_workers=5) as executor:
        futures = [executor.submit(worker, i) for i in range(1, 10)]
        for f in concurrent.futures.as_completed(futures):
            f.result() # Will raise if any thread failed with lock or corruption error

    # Re-open final file to confirm valid structure
    org_wb = openpyxl.load_workbook(tmp_path / eh.ORGANIZED)
    assert org_wb["Organized Output"].max_row == 11


def test_should_stop_allows_clean_interruption(monkeypatch, tmp_path):
    from backend import pipeline
    rows = [
        {"row": 2, "mpn": "STOP1", "mfr": "Brand"},
        {"row": 3, "mpn": "STOP2", "mfr": "Brand"},
        {"row": 4, "mpn": "STOP3", "mfr": "Brand"},
    ]
    results = {}
    ckpt = tmp_path / "ckpt.json"

    monkeypatch.setattr(pipeline, "process_row", lambda r, *a, **k: {
        "status": "SUCCESS", "tier": "1", "sources": [], "score": 90,
        "confidence": "HIGH", "product_title": f"P {r['mpn']}"
    })

    export_log = []
    # Stop condition: stop after 1 row has been processed
    cfg = {
        "save_every": 1,
        "on_export": lambda r, res, is_final=False: export_log.append((len(res), is_final)),
        "should_stop": lambda: len(results) >= 1
    }
    res = pipeline.run_batch(rows, results, ckpt, cfg)

    # Cleanly stopped after 1 row
    assert len(res) == 1
    assert 2 in res
    assert 3 not in res
    # Final export was still triggered
    assert any(is_final for _, is_final in export_log)




