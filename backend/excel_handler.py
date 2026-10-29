"""Non-destructive workbook handling + atomic JSON checkpoints."""
import hashlib, json, os
from urllib.parse import quote_plus
import openpyxl
from config import settings

ALIASES = {"mpn": ["mpn1", "mpn", "part number"], "mfr": ["manufacturer", "brand"],
           "title": ["db_name", "product name", "title"], "url": ["vserve_source_url", "data_url1"]}


def sheet_names(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    names = wb.sheetnames; wb.close(); return names


def headers(ws):
    return {str(c.value).strip().lower(): c.column for c in ws[1] if c.value is not None}


def detect_columns(ws):
    h = headers(ws)
    return {k: next((a for a in al if a in h), None) for k, al in ALIASES.items()}


def read_rows(ws, cmap, limit=None):
    h, rows = headers(ws), []
    def val(r, key):
        return "" if not cmap.get(key) or cmap[key] not in h else str(ws.cell(r, h[cmap[key]]).value or "").strip()
    for r in range(2, ws.max_row + 1):
        row = {"row": r, "mpn": val(r, "mpn"), "mfr": val(r, "mfr"), "title": val(r, "title"), "url": val(r, "url")}
        if row["mpn"] and row["mfr"]:
            rows.append(row)
    seen, uniq = set(), []
    for x in rows:                      # drop duplicate MPN+Manufacturer rows
        k = (x["mpn"].lower(), x["mfr"].lower())
        if k not in seen:
            seen.add(k); uniq.append(x)
    return uniq[:limit] if limit else uniq


def _col(ws, name):
    h = headers(ws)
    if name.lower() in h:
        return h[name.lower()]
    c = ws.max_column + 1
    ws.cell(1, c, name)
    return c


def write_result(ws, row, res):
    def put(name, v):
        if v not in (None, ""):
            ws.cell(row, _col(ws, name), v)
    put("Image Link", res.get("image"))
    for i, u in enumerate(res.get("sources", [])[:3], 1):
        put(f"Source {i}", u)
    put("Google Search Link", res.get("google"))
    ov = res.get("overview", "")
    put("Overview", ov)
    for i in range(5):
        put(f"overview_{i+1}", ov[i*1000:(i+1)*1000])
    feats = res.get("features", [])
    put("Features", "\n".join(feats))
    for i, f in enumerate(feats[:35], 1):
        put(f"feature_{i}", f)
    specs = res.get("specifications", [])
    put("Specifications", "\n".join(f'{s["name"]}: {s["value"]}' for s in specs))
    for i, s in enumerate(specs[:35], 1):
        put(f"specification_header{i}", s["name"]); put(f"specification_{i}", s["value"])
    for i, f in enumerate(res.get("fitments", [])[:60], 1):
        put(f"other_feature_{i}", f)
    ok = res["status"] in ("SUCCESS", "REVIEW")
    put("Completed", "Y" if ok else "N"); put("DataNotAvailable", "N" if ok else "Y")
    put("reworkreason", res.get("reason", "") if res["status"] != "SUCCESS" else "")


def google_link(mfr, mpn):
    return "https://www.google.com/search?q=" + quote_plus(f"{mfr} {mpn}")


def ckpt_path(filename, size):
    return settings.CACHE_DIR / f"checkpoint_{hashlib.md5(f'{filename}{size}'.encode()).hexdigest()}.json"


def load_ckpt(p):
    return {int(k): v for k, v in json.loads(p.read_text()).items()} if p.exists() else {}


def save_ckpt(p, results):
    tmp = p.with_suffix(".tmp")
    tmp.write_text(json.dumps(results))
    os.replace(tmp, p)
