import json, os, threading, time
from pathlib import Path
import openpyxl
from config import settings
from backend import excel_handler as xh

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRAPED, REVIEW = "Data_collection_scraped.xlsx", "Scrape_Review_Report.xlsx"
ORGANIZED = "Final_Organized_Output.xlsx"
FILL = {"SUCCESS": "C6EFCE", "REVIEW": "FFEB9C", "FAILED": "FFC7CE", "PENDING": "E0E0E0"}
COLS = [("Row", 6), ("MPN", 18), ("Manufacturer", 18), ("Product Name", 36), ("Status", 11), ("Confidence", 12),
        ("Image", 12), ("Source 1", 38), ("Source 2", 38), ("Source 3", 38), ("Overview", 60),
        ("Features", 60), ("Specifications", 60), ("Fitments", 45), ("Missing / Review Reason", 32),
        ("Google Search Link", 38)]

_EXPORT_LOCK = threading.Lock()


def _safe_save_wb(wb, target_path, max_retries=5, base_delay=0.1):
    target_path = Path(target_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = target_path.with_name(f"{target_path.name}.tmp")
    for attempt in range(max_retries):
        try:
            wb.save(tmp_path)
            os.replace(tmp_path, target_path)
            return
        except (PermissionError, OSError) as e:
            if attempt == max_retries - 1:
                try:
                    if tmp_path.exists():
                        tmp_path.unlink()
                except Exception:
                    pass
                raise e
            time.sleep(base_delay * (2 ** attempt))


def _safe_write_text(target_path, content, max_retries=5, base_delay=0.1):
    target_path = Path(target_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = target_path.with_name(f"{target_path.name}.tmp")
    for attempt in range(max_retries):
        try:
            tmp_path.write_text(content, encoding="utf-8")
            os.replace(tmp_path, target_path)
            return
        except (PermissionError, OSError) as e:
            if attempt == max_retries - 1:
                try:
                    if tmp_path.exists():
                        tmp_path.unlink()
                except Exception:
                    pass
                raise e
            time.sleep(base_delay * (2 ** attempt))


def organized(rows, results):
    """One tidy, formatted row per product in the original input order."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Organized Output"
    ws.append([c for c, _ in COLS])
    seen_rows = set()
    for r in rows:
        row_id = r["row"]
        if row_id in seen_rows:
            continue
        seen_rows.add(row_id)
        x = results.get(row_id)
        if not x:
            glink = xh.google_link(r["mfr"], r["mpn"])
            ws.append([row_id, r["mpn"], r["mfr"], "", "PENDING", "",
                       "", "", "", "", "", "", "", "", "Pending processing", glink])
            i = ws.max_row
            if glink:
                ws.cell(i, 16).hyperlink = glink
                ws.cell(i, 16).font = Font(color="0563C1", underline="single")
            ws.cell(i, 5).fill = PatternFill("solid", fgColor=FILL["PENDING"])
            continue

        src = (x.get("sources") or []) + ["", "", ""]
        ws.append([row_id, r["mpn"], r["mfr"], x.get("product_title", ""), x["status"], x.get("confidence", ""),
                   "View" if x.get("image") else "", src[0], src[1], src[2],
                   (x.get("overview") or "")[:32000],
                   "\n".join(f"• {f}" for f in x.get("features", []))[:32000],
                   "\n".join(f'{s["name"]}: {s["value"]}' for s in x.get("specifications", []))[:32000],
                   "\n".join(x.get("fitments", []))[:32000], x.get("reason", ""), x.get("google", "")])
        i = ws.max_row
        if x.get("image"):
            ws.cell(i, 7).hyperlink = x["image"]; ws.cell(i, 7).font = Font(color="0563C1", underline="single")
        for c in (8, 9, 10, 16):
            if ws.cell(i, c).value:
                ws.cell(i, c).hyperlink = ws.cell(i, c).value
                ws.cell(i, c).font = Font(color="0563C1", underline="single")
        ws.cell(i, 5).fill = PatternFill("solid", fgColor=FILL.get(x["status"], "FFFFFF"))
    for n, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(n)].width = w
        ws.cell(1, n).fill = PatternFill("solid", fgColor="0F1B3D")
        ws.cell(1, n).font = Font(bold=True, color="FFFFFF")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0].row].height = 120
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    _safe_save_wb(wb, settings.OUTPUT_DIR / ORGANIZED)


def export(src_path, sheet, rows, results, is_final=None):
    with _EXPORT_LOCK:
        src_path = Path(src_path)
        wb = openpyxl.load_workbook(src_path)
        ws = wb[sheet]
        seen_rows = set()
        for r in rows:
            row_id = r["row"]
            if row_id in seen_rows:
                continue
            seen_rows.add(row_id)
            if row_id in results:
                xh.write_result(ws, row_id, results[row_id])
        out = settings.OUTPUT_DIR / SCRAPED
        _safe_save_wb(wb, out)
        organized(rows, results)
        rp = openpyxl.Workbook()
        s = rp.active; s.title = "Summary"
        stat = [x["status"] for x in results.values()]
        total = len(rows)
        from urllib.parse import urlparse
        doms = {urlparse(u).netloc for x in results.values() for u in x.get("sources", [])}
        pending_count = sum(1 for r in rows if r["row"] not in results)
        summary_items = [("Total rows", total), ("Verified (SUCCESS)", stat.count("SUCCESS")),
                         ("Review required", stat.count("REVIEW")), ("Failed", stat.count("FAILED"))]
        if pending_count > 0:
            summary_items.append(("Pending", pending_count))
        summary_items.append(("Distinct source websites", len(doms)))
        for k, v in summary_items:
            s.append([k, v])
        d = rp.create_sheet("Row Audit")
        d.append(["Row", "MPN", "Manufacturer", "Tier", "Selected URL", "Score", "Confidence", "Status",
                  "Evidence", "Missing / Reason", "Recommended action"])
        seen_rows = set()
        for r in rows:
            row_id = r["row"]
            if row_id in seen_rows:
                continue
            seen_rows.add(row_id)
            x = results.get(row_id)
            if not x:
                glink = xh.google_link(r["mfr"], r["mpn"])
                d.append([row_id, r["mpn"], r["mfr"], "", "", "",
                          "", "PENDING", "", "Pending processing", f"Pending: {glink}"])
                continue
            g_link = x.get("google") or xh.google_link(r["mfr"], r["mpn"])
            act = "None" if x["status"] == "SUCCESS" else f'Manual check: {g_link}'
            d.append([row_id, r["mpn"], r["mfr"], x.get("tier"), (x["sources"] or [""])[0], x.get("score"),
                      x.get("confidence"), x["status"], "; ".join(x.get("notes", [])), x.get("reason", ""), act])
        _safe_save_wb(rp, settings.OUTPUT_DIR / REVIEW)

        should_write_history = is_final is True or (is_final is None and len(results) >= total)
        if should_write_history:
            hist = settings.OUTPUT_DIR / "history.json"
            h = json.loads(hist.read_text()) if hist.exists() else []
            is_partial = len(results) < total
            rec = {
                "time": time.strftime("%Y-%m-%d %H:%M:%S"),
                "file": src_path.name,
                "total": total,
                "success": stat.count("SUCCESS"),
                "review": stat.count("REVIEW"),
                "failed": stat.count("FAILED"),
                "status": "PARTIAL" if is_partial else "COMPLETE"
            }
            h.append(rec)
            _safe_write_text(hist, json.dumps(h, indent=2))
        return out, settings.OUTPUT_DIR / REVIEW
