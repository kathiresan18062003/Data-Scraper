import json
import os
import re
import threading
import time
from pathlib import Path
from urllib.parse import quote_plus, urlparse
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.app.config import OUTPUT_DIR
from backend.app.schemas.job import SheetInfo

ALIASES = {
    "sno": ["s.no", "sno", "sl.no", "sl no", "serial no", "serial number", "id", "row", "#"],
    "mpn1": ["mpn1", "mpn 1", "part number 1", "model 1", "primary mpn"],
    "mpn": ["mpn", "part number", "part#", "item number", "item#", "sku", "model"],
    "mfr": ["manufacturer", "brand", "mfr", "make"],
    "title": ["product name", "product_name", "title", "name", "db_name", "description"],
    "upc": ["upc", "upc code", "barcode", "ean"],
    "url": ["vserve_source_url", "data_url1", "source_url", "url", "link"]
}

FILL = {
    "SUCCESS": "C6EFCE",
    "REVIEW": "FFEB9C",
    "FAILED": "FFC7CE",
    "PENDING": "E0E0E0"
}

COLS = [
    ("Row", 6), ("MPN", 18), ("Manufacturer", 18), ("Product Name", 36), ("Status", 11), ("Confidence", 12),
    ("Image", 12), ("Source 1", 38), ("Source 2", 38), ("Source 3", 38), ("Overview", 60),
    ("Features", 60), ("Specifications", 60), ("Fitments", 45), ("Missing / Review Reason", 32),
    ("Google Search Link", 38)
]

_EXPORT_LOCK = threading.Lock()


def google_link(mfr: str, mpn: str) -> str:
    return "https://www.google.com/search?q=" + quote_plus(f"{mfr} {mpn}".strip())


def checkpoint_path(job_id: str) -> Path:
    return OUTPUT_DIR / f"{job_id}_checkpoint.json"


def save_checkpoint(job_id: str, data: dict):
    cp = checkpoint_path(job_id)
    cp.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = cp.with_suffix(".tmp")
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, default=str)
    try:
        os.replace(tmp_path, cp)
    except OSError:
        # Fallback for Windows lock contention
        time.sleep(0.1)
        os.replace(tmp_path, cp)


def load_checkpoint(job_id: str) -> dict | None:
    cp = checkpoint_path(job_id)
    if cp.exists():
        try:
            with open(cp, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None
    return None


def _safe_save_wb(wb: openpyxl.Workbook, target_path: Path, max_retries: int = 5, base_delay: float = 0.1):
    target_path = Path(target_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = target_path.with_name(f"{target_path.stem}_{os.getpid()}_{int(time.time()*1000)}.tmp")
    for attempt in range(max_retries):
        try:
            wb.save(tmp_path)
            if target_path.exists():
                try:
                    os.replace(tmp_path, target_path)
                except (PermissionError, OSError):
                    time.sleep(base_delay * (2 ** attempt))
                    os.replace(tmp_path, target_path)
            else:
                os.replace(tmp_path, target_path)
            return
        except (PermissionError, OSError) as e:
            if attempt == max_retries - 1:
                if tmp_path.exists():
                    try:
                        tmp_path.unlink()
                    except Exception:
                        pass
                raise e
            time.sleep(base_delay * (2 ** attempt))


def sheet_names(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def get_headers(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, int]:
    for row_idx in (1, 2):
        row_cells = [c for c in ws[row_idx] if c.value is not None and str(c.value).strip()]
        if row_cells:
            return {str(c.value).strip().lower(): c.column for c in row_cells}
    return {}


def detect_columns(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, str | None]:
    h = get_headers(ws)
    detected = {}
    for key, aliases in ALIASES.items():
        found = None
        for a in aliases:
            for header_name in h.keys():
                if a == header_name or a in header_name:
                    found = header_name
                    break
            if found:
                break
        detected[key] = found
    return detected


def get_sheet_info(file_path: Path) -> list[SheetInfo]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers_list = [str(c.value).strip() for c in ws[1] if c.value is not None and str(c.value).strip()]
        if not headers_list and ws.max_row >= 2:
            headers_list = [str(c.value).strip() for c in ws[2] if c.value is not None and str(c.value).strip()]
        detected = detect_columns(ws)

        preview_rows = []
        for row in list(ws.iter_rows(values_only=True))[:11]:
            preview_rows.append([str(c) if c is not None else "" for c in row])

        sheets.append(SheetInfo(
            name=sheet_name,
            row_count=max(ws.max_row - 1, 0),
            column_count=ws.max_column or 0,
            detected_columns=detected,
            headers=headers_list,
            preview_rows=preview_rows
        ))
    wb.close()
    return sheets


def read_rows(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    cmap: dict[str, str],
    limit: int | None = None,
    start_row: int | None = None,
    end_row: int | None = None,
    sno_start: int | None = None,
    sno_end: int | None = None,
) -> list[dict]:
    h = get_headers(ws)

    def val(r: int, key: str) -> str:
        col_name = (cmap.get(key) or "").strip().lower()
        if not col_name:
            return ""
        col_idx = h.get(col_name)
        if not col_idx:
            for header, idx in h.items():
                if header == col_name or col_name in header or header in col_name:
                    col_idx = idx
                    break
        if not col_idx:
            return ""
        cell_val = ws.cell(r, col_idx).value
        return str(cell_val).strip() if cell_val is not None else ""

    rows = []
    min_r = max(start_row, 2) if start_row else 2
    max_r = min(end_row, ws.max_row) if end_row else ws.max_row

    for r in range(min_r, max_r + 1):
        sno = val(r, "sno") or str(r - 1)
        mpn1 = val(r, "mpn1")
        mpn = val(r, "mpn")
        primary_mpn = mpn1 if mpn1 else mpn
        mfr = val(r, "mfr")
        title = val(r, "title")
        upc = val(r, "upc")
        url = val(r, "url")

        # Range filtering by S.No if provided
        if sno_start is not None or sno_end is not None:
            try:
                sno_clean = re.sub(r"[^\d]", "", sno)
                if sno_clean:
                    sno_num = int(sno_clean)
                    if sno_start is not None and sno_num < sno_start:
                        continue
                    if sno_end is not None and sno_num > sno_end:
                        continue
            except (ValueError, TypeError):
                pass

        if primary_mpn:
            rows.append({
                "row": r,
                "sno": sno,
                "mpn": primary_mpn,
                "mpn1": mpn1,
                "base_mpn": mpn,
                "mfr": mfr or "Unknown",
                "title": title,
                "upc": upc,
                "url": url,
            })

    seen = set()
    uniq = []
    for x in rows:
        k = (x["mpn"].lower(), x["mfr"].lower())
        if k not in seen:
            seen.add(k)
            uniq.append(x)

    return uniq[:limit] if limit else uniq


def _col(ws, name: str) -> int:
    h = get_headers(ws)
    if name.lower() in h:
        return h[name.lower()]
    c = ws.max_column + 1
    ws.cell(1, c, name)
    return c


def write_result_to_worksheet(ws, row_idx: int, res: dict):
    def put(name: str, v):
        if v not in (None, ""):
            ws.cell(row_idx, _col(ws, name), v)

    put("Image Link", res.get("image") or res.get("image_url"))
    sources = res.get("sources") or res.get("source_urls") or []
    for i, u in enumerate(sources[:3], 1):
        put(f"Source {i}", u)
    put("Google Search Link", res.get("google") or res.get("google_search_url"))
    ov = res.get("overview", "")
    put("Overview", ov)
    for i in range(5):
        put(f"overview_{i+1}", ov[i*1000:(i+1)*1000])

    feats = res.get("features", [])
    put("Features", "\n".join(feats))
    for i, f in enumerate(feats[:35], 1):
        put(f"feature_{i}", f)

    specs = res.get("specifications", [])
    put("Specifications", "\n".join(f'{s.get("name")}: {s.get("value")}' for s in specs))
    for i, s in enumerate(specs[:35], 1):
        put(f"specification_header{i}", s.get("name"))
        put(f"specification_{i}", s.get("value"))

    fitments = res.get("fitments", [])
    for i, f in enumerate(fitments[:60], 1):
        put(f"other_feature_{i}", f)

    ok = res.get("status") in ("SUCCESS", "REVIEW")
    put("Completed", "Y" if ok else "N")
    put("DataNotAvailable", "N" if ok else "Y")
    put("reworkreason", res.get("reason", "") or res.get("rework_reason", "") if not ok else "")


def export_organized_workbook(rows: list[dict], results: dict[int, dict], target_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Organized Output"
    ws.append([c for c, _ in COLS])

    seen_rows = set()
    for r in rows:
        row_id = r["row"]
        if row_id in seen_rows:
            continue
        seen_rows.add(row_id)
        x = results.get(row_id)
        glink = (x.get("google") or x.get("google_search_url") or google_link(r["mfr"], r["mpn"])) if x else google_link(r["mfr"], r["mpn"])

        if not x:
            ws.append([
                row_id, r["mpn"], r["mfr"], "", "PENDING", "",
                "", "", "", "", "", "", "", "", "Pending processing", glink
            ])
            i = ws.max_row
            if glink:
                ws.cell(i, 16).hyperlink = glink
                ws.cell(i, 16).font = Font(color="0563C1", underline="single")
            ws.cell(i, 5).fill = PatternFill("solid", fgColor=FILL["PENDING"])
            continue

        src = (x.get("sources") or x.get("source_urls") or []) + ["", "", ""]
        img = x.get("image") or x.get("image_url") or ""
        features = x.get("features") or []
        specs = x.get("specifications") or []
        fitments = x.get("fitments") or []
        reason = x.get("reason") or x.get("rework_reason") or ""

        ws.append([
            row_id, r["mpn"], r["mfr"], x.get("product_title") or x.get("product_name") or "",
            x.get("status", "PENDING"), x.get("confidence", ""),
            "View" if img else "", src[0], src[1], src[2],
            (x.get("overview") or "")[:32000],
            "\n".join(f"• {f}" for f in features)[:32000],
            "\n".join(f'{s.get("name")}: {s.get("value")}' for s in specs)[:32000],
            "\n".join(fitments)[:32000],
            reason, glink
        ])
        i = ws.max_row
        if img:
            ws.cell(i, 7).hyperlink = img
            ws.cell(i, 7).font = Font(color="0563C1", underline="single")
        for c in (8, 9, 10, 16):
            if ws.cell(i, c).value:
                ws.cell(i, c).hyperlink = str(ws.cell(i, c).value)
                ws.cell(i, c).font = Font(color="0563C1", underline="single")
        ws.cell(i, 5).fill = PatternFill("solid", fgColor=FILL.get(x.get("status"), "FFFFFF"))

    for n, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(n)].width = w
        ws.cell(1, n).fill = PatternFill("solid", fgColor="0F1B3D")
        ws.cell(1, n).font = Font(bold=True, color="FFFFFF")

    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0].row].height = 100

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    _safe_save_wb(wb, target_path)


def export_review_workbook(rows: list[dict], results: dict[int, dict], target_path: Path):
    rp = openpyxl.Workbook()
    s = rp.active
    s.title = "Summary"

    stat = [x.get("status") for x in results.values()]
    total = len(rows)
    doms = {
        urlparse(u).netloc
        for x in results.values()
        for u in (x.get("sources") or x.get("source_urls") or [])
        if u
    }
    pending_count = sum(1 for r in rows if r["row"] not in results)

    summary_items = [
        ("Total rows", total),
        ("Verified (SUCCESS)", stat.count("SUCCESS")),
        ("Review required", stat.count("REVIEW")),
        ("Failed", stat.count("FAILED"))
    ]
    if pending_count > 0:
        summary_items.append(("Pending", pending_count))
    summary_items.append(("Distinct source websites", len(doms)))

    for k, v in summary_items:
        s.append([k, v])

    d = rp.create_sheet("Row Audit")
    d.append([
        "Row", "MPN", "Manufacturer", "Tier", "Selected URL", "Score", "Confidence", "Status",
        "Evidence", "Missing / Reason", "Recommended action"
    ])

    seen_rows = set()
    for r in rows:
        row_id = r["row"]
        if row_id in seen_rows:
            continue
        seen_rows.add(row_id)
        x = results.get(row_id)
        g_link = (x.get("google") or x.get("google_search_url") or google_link(r["mfr"], r["mpn"])) if x else google_link(r["mfr"], r["mpn"])

        if not x:
            d.append([
                row_id, r["mpn"], r["mfr"], "", "", "", "", "PENDING", "", "Pending processing", f"Pending: {g_link}"
            ])
            continue

        act = "None" if x.get("status") == "SUCCESS" else f"Manual check: {g_link}"
        sources = x.get("sources") or x.get("source_urls") or []
        notes = x.get("notes") or x.get("validation_notes") or []
        reason = x.get("reason") or x.get("rework_reason") or ""

        d.append([
            row_id, r["mpn"], r["mfr"], x.get("tier"), sources[0] if sources else "", x.get("score"),
            x.get("confidence"), x.get("status"), "; ".join(notes), reason, act
        ])

    _safe_save_wb(rp, target_path)


def export_all(src_path: Path, sheet: str, rows: list[dict], results: dict[int, dict], job_id: str | None = None) -> dict[str, Path]:
    with _EXPORT_LOCK:
        src_path = Path(src_path)
        wb = openpyxl.load_workbook(src_path)
        ws = wb[sheet]

        seen_rows = set()
        for r in rows:
            row_id = r["row"]
            if row_id in seen_rows:
                continue
            seen_rows.add(row_id)
            if row_id in results:
                write_result_to_worksheet(ws, row_id, results[row_id])

        # Paths
        prefix = f"{job_id}_" if job_id else ""
        out_scraped = OUTPUT_DIR / f"{prefix}Data_collection_scraped.xlsx"
        out_organized = OUTPUT_DIR / f"{prefix}Final_Organized_Output.xlsx"
        out_review = OUTPUT_DIR / f"{prefix}Scrape_Review_Report.xlsx"

        _safe_save_wb(wb, out_scraped)
        export_organized_workbook(rows, results, out_organized)
        export_review_workbook(rows, results, out_review)

        # Also update root output files for quick convenience
        if job_id:
            try:
                _safe_save_wb(wb, OUTPUT_DIR / "Data_collection_scraped.xlsx")
                export_organized_workbook(rows, results, OUTPUT_DIR / "Final_Organized_Output.xlsx")
                export_review_workbook(rows, results, OUTPUT_DIR / "Scrape_Review_Report.xlsx")
            except Exception:
                pass

        return {
            "scraped": out_scraped,
            "organized": out_organized,
            "review": out_review
        }
